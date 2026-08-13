"""PPTX → KBDoc.

Slide decks arrive as training material and committee packs. The slide is the natural page:
its title is the heading, its body text and tables sit under it, and speaker notes are kept
because they often carry the actual policy statement the slide only alludes to.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from kb_schemas.kbdoc import KBDoc
from pptx import Presentation

from kb_idp.builder import KBDocBuilder


def parse_pptx(data: bytes) -> KBDoc:
    presentation = Presentation(BytesIO(data))
    builder = KBDocBuilder(source_format="pptx")
    builder.note_engine("python-pptx", "native")

    for page, slide in enumerate(presentation.slides, start=1):
        title = _slide_title(slide)
        if title:
            builder.add_heading(title, page=page)
        else:
            builder.add_heading(f"Slide {page}", page=page)

        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_table:
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                builder.add_table(rows, header_rows=1, page=page)
            elif shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs)
                    builder.add(text, page=page)

        notes = _speaker_notes(slide)
        if notes:
            builder.add(notes, page=page)

    return builder.build(page_count=len(presentation.slides), source_format="pptx")


def _slide_title(slide: Any) -> str:
    title_shape = slide.shapes.title
    if title_shape is not None and title_shape.has_text_frame:
        return str(title_shape.text_frame.text).strip()
    return ""


def _speaker_notes(slide: Any) -> str:
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    return str(frame.text).strip() if frame is not None else ""
