"""DOCX → KBDoc.

Body order matters: a table between two paragraphs belongs between them, and python-docx's
`document.paragraphs` / `document.tables` return them separately. Walking the XML body keeps
the document's actual reading order, which the section tracker depends on.
"""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from kb_schemas.kbdoc import KBDoc

from kb_idp.builder import KBDocBuilder

#: Word styles that mean "heading" regardless of the text. Vietnamese Word installs localize
#: the style name, so both are matched.
_HEADING_PREFIXES = ("heading", "title", "đầu đề", "tiêu đề")


def parse_docx(data: bytes) -> KBDoc:
    document: DocxDocument = Document(BytesIO(data))
    builder = KBDocBuilder(source_format="docx")
    builder.note_engine("python-docx", "native")

    core_title = (document.core_properties.title or "").strip()
    if core_title:
        builder.set_title_hint(core_title)

    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)
            text = paragraph.text
            if not text.strip():
                continue
            builder.add(
                text,
                block_type="heading" if _is_heading(paragraph) else None,
                page=1,  # DOCX has no fixed pagination; the review editor uses block order
            )
        elif element.tag == qn("w:tbl"):
            table = DocxTable(element, document)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            builder.add_table(rows, header_rows=1 if _looks_like_header(rows) else 0, page=1)

    return builder.build(page_count=1, source_format="docx")


def _is_heading(paragraph: Paragraph) -> bool:
    style = (paragraph.style.name or "").lower() if paragraph.style else ""
    if any(style.startswith(prefix) for prefix in _HEADING_PREFIXES):
        return True
    # Unstyled documents are the norm in this corpus — a short, fully bold paragraph is the
    # other reliable signal, and the section tracker catches "Điều 12" regardless of styling.
    runs = [run for run in paragraph.runs if run.text.strip()]
    return bool(runs) and all(run.bold for run in runs) and len(paragraph.text) < 120


def _looks_like_header(rows: list[list[str]]) -> bool:
    if len(rows) < 2:
        return False
    first, second = rows[0], rows[1]
    # A header row is text where the data rows hold numbers.
    return (
        any(cell.strip() for cell in first)
        and sum(
            1
            for cell in second
            if cell.replace(",", "").replace(".", "").replace("%", "").isdigit()
        )
        > 0
    )
