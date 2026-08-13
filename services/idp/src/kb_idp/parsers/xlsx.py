"""XLSX → KBDoc.

Spreadsheets in this corpus are fee schedules, limit tables and risk-weight matrices: the
answer to a question is a *cell*, and it is only meaningful with its row and column headers.
So each sheet becomes one table block with the header row marked, rather than a stream of
loose values.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from kb_schemas.kbdoc import KBDoc
from openpyxl import load_workbook

from kb_idp.builder import KBDocBuilder

#: Guard against a stray formatted-to-row-1048576 sheet producing a million empty rows.
_MAX_ROWS = 5000
_MAX_COLS = 100


def parse_xlsx(data: bytes) -> KBDoc:
    workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    builder = KBDocBuilder(source_format="xlsx")
    builder.note_engine("openpyxl", "native")

    try:
        for index, sheet in enumerate(workbook.worksheets, start=1):
            if sheet.sheet_state != "visible":
                # Hidden sheets are usually working calculations, not published content.
                builder.warn(f"skipped hidden sheet {sheet.title!r}")
                continue

            builder.add_heading(sheet.title, page=index)
            rows = _read_rows(sheet)
            if not rows:
                builder.warn(f"sheet {sheet.title!r} is empty")
                continue
            builder.add_table(rows, header_rows=1, page=index)
    finally:
        workbook.close()

    return builder.build(page_count=len(workbook.worksheets), source_format="xlsx")


def _read_rows(sheet: Any) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in sheet.iter_rows(max_row=_MAX_ROWS, max_col=_MAX_COLS, values_only=True):
        cells = ["" if value is None else str(value).strip() for value in row]
        while cells and not cells[-1]:
            cells.pop()
        if cells:
            rows.append(cells)
    width = max((len(row) for row in rows), default=0)
    return [row + [""] * (width - len(row)) for row in rows]
